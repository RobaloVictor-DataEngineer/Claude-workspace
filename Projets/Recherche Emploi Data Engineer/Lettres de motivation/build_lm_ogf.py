#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Construit la lettre de motivation OGF (Ingénieur Data, Courbevoie) à partir du modèle.
Adresse : Île-de-France (règle du profil, offre à Courbevoie)."""
import docx

SRC = "_MODELE_lettre_motivation.docx"
OUT = "LM_ROBALO_Victor_OGF_IngenieurData.docx"

d = docx.Document(SRC)
paras = d.paragraphs

def set_run(p, i, text):
    p.runs[i].text = text

# --- Para 1 : date (pas de mention de ville, cf règle Île-de-France -> pas de rue) ---
p = paras[1]
set_run(p, 2, "")
set_run(p, 3, "07")
set_run(p, 4, "/0")
set_run(p, 5, "9")
set_run(p, 6, "/2026")

# --- Para 2 : adresse -> Mobile Île-de-France ---
p = paras[2]
set_run(p, 1, " Mobile Île-de-France")

# --- Para 3 : ligne code postal / ville -> vide ---
p = paras[3]
set_run(p, 0, "")

# --- Para 6 : nom entreprise (bandeau droit) ---
p = paras[6]
set_run(p, 3, "Groupe OGF")

# --- Para 7 : intitulé du poste (bandeau droit) ---
p = paras[7]
set_run(p, 0, "\t               Ingénieur Data")

# --- Para 13 : objet ---
p = paras[13]
set_run(p, 0, "Objet : Candidature au poste d'Ingénieur Data")

# --- Para 15 : accroche ---
p = paras[15]
set_run(p, 0, "C'est avec un vif intérêt que je vous transmets ma candidature pour le poste d'Ingénieur Data. Garantir la mise à disposition de données fiables, documentées et disponibles dans les délais, au service des équipes métiers et de la Data Viz, est précisément le cœur de mon métier depuis trois ans.")
set_run(p, 1, "")
set_run(p, 2, "")

# --- Para 16 : expérience Sanofi ---
p = paras[16]
set_run(p, 0, "Mon expérience de trois ans en alternance chez Sanofi m'a appris ce que signifie fiabiliser un flux de données en production : j'y ai transformé des données brutes en indicateurs de performance (KPIs) exploités par les équipes métier et le Comité de Direction, dans un environnement industriel exigeant sur la qualité, la traçabilité et la gestion des anomalies.")

# --- Para 17 : expérience La Royale Gaming Investment ---
p = paras[17]
set_run(p, 0, "Mon stage de fin d'études chez ")
set_run(p, 1, "La Royale Gaming Investment")
set_run(p, 2, ", à Lisbonne, m'a permis de construire et maintenir des flux ETL (dbt/Dataform) sur BigQuery, avec supervision quotidienne des traitements et optimisation des requêtes et des coûts de 15 à 20 %. J'y ai aussi géré la gouvernance des données avec les équipes métiers et alimenté des Data Marts pour des dashboards Power BI, sur des volumes de transactions financières significatifs.")

# --- Para 18 : compétences techniques (honnête sur les manques) ---
p = paras[18]
set_run(p, 0, "Sur le plan technique, je pratique SQL avancé au quotidien (requêtes complexes, optimisation de performance) ainsi que SQL Server, Snowflake et BigQuery pour la modélisation d'architectures décisionnelles (Data Warehouse, Data Marts). Je n'ai en revanche pas encore d'expérience opérationnelle sur Talend ni sur l'écosystème Azure : mon expérience ETL s'est construite avec dbt/Dataform, mais la logique de construction de flux, d'industrialisation et de gestion des rejets est transposable, et je suis motivé pour monter rapidement en compétence sur ces outils.")

# --- Para 19 : savoir-être ---
p = paras[19]
set_run(p, 0, "Pragmatique et rigoureux, j'apprécie les environnements où la fiabilité des données en production est une priorité partagée entre équipes techniques et métiers, et où le Build s'articule étroitement avec le Run.")

# --- Para 20 : clôture ---
p = paras[20]
set_run(p, 0, "Je serais ravi d'échanger avec vous sur la manière dont mon profil peut contribuer à la fiabilisation et à l'évolution de votre plateforme data.")

# Para 21 (formule de politesse) : inchangée
# Para 26 (signature) : inchangée

# --- Nettoyage du pied de page (résidu d'une candidature précédente) ---
for s in d.sections:
    for p in s.footer.paragraphs:
        for r in p.runs:
            r.text = ""

d.save(OUT)
print("LM générée :", OUT)

d2 = docx.Document(OUT)
print("\n=== CONTENU FINAL ===")
for i, p in enumerate(d2.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")
print("\n=== FOOTER ===")
for s in d2.sections:
    for i, p in enumerate(s.footer.paragraphs):
        print(i, repr(p.text))
