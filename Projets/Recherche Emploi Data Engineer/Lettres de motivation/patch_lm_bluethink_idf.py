#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Patch LM_ROBALO_Victor_Bluethink_DataEngineer.docx in place:
remove Rouen mentions (street + city + date-line Rouen) and put Île-de-France,
same pattern as the Aneo letter."""
import docx

F = "LM_ROBALO_Victor_Bluethink_DataEngineer.docx"
d = docx.Document(F)
paras = d.paragraphs

# --- Para 1 : date line - remove "Rouen, " ---
p = paras[1]
p.runs[2].text = ""

# --- Para 2 : address block first line -> Mobile Île-de-France ---
p = paras[2]
p.runs[1].text = " Mobile Île-de-France"

# --- Para 3 : postal code / city line -> blank ---
p = paras[3]
p.runs[0].text = ""

d.save(F)
print("LM patchée (Île-de-France) :", F)

d2 = docx.Document(F)
print("\n=== CONTENU FINAL ===")
for i, p in enumerate(d2.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")
