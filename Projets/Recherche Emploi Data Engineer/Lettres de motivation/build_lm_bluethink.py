#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Construit la lettre de motivation Bluethink (Data Engineer) à partir du modèle,
en ne touchant qu'au texte des runs concernés — police/taille/gras/alignement du
modèle sont conservés intégralement."""
import docx

SRC = "_MODELE_lettre_motivation.docx"
OUT = "LM_ROBALO_Victor_Bluethink_DataEngineer.docx"

d = docx.Document(SRC)
paras = d.paragraphs

def set_run(p, i, text):
    p.runs[i].text = text

# --- Para 1 : date ---
p = paras[1]
set_run(p, 2, "Rouen, ")
set_run(p, 3, "07")
set_run(p, 4, "/0")
set_run(p, 5, "9")
set_run(p, 6, "/2026")

# --- Para 6 : nom entreprise (bandeau droit) ---
p = paras[6]
set_run(p, 3, "Bluethink")

# --- Para 7 : intitulé du poste (bandeau droit) ---
p = paras[7]
set_run(p, 0, "\t               Data Engineer")

# --- Para 13 : objet ---
p = paras[13]
set_run(p, 0, "Objet : Candidature au poste de Data Engineer")

# --- Para 15 : accroche ---
p = paras[15]
set_run(p, 0, "C'est avec un vif intérêt que je vous transmets ma candidature pour le poste de Data Engineer, au sein d'une équipe Data à taille humaine qui construit une plateforme centralisant billetterie, CRM et outils marketing pour mieux comprendre et accompagner ses clients. Travailler en lien direct avec la Head of Data, sur une plateforme qui alimente concrètement des produits, est exactement le type de contexte où j'ai envie d'apporter ma contribution.")
set_run(p, 1, "")
set_run(p, 2, "")

# --- Para 16 : expérience Sanofi ---
p = paras[16]
set_run(p, 0, "Mon expérience de trois ans en alternance chez Sanofi m'a appris ce que signifie fiabiliser un flux de données en production : j'y ai transformé des données brutes en indicateurs de performance (KPIs) exploités par les équipes métier et le Comité de Direction, dans un environnement industriel exigeant sur la qualité et la traçabilité.")

# --- Para 17 : expérience La Royale Gaming Investment ---
p = paras[17]
set_run(p, 0, "Mon stage de fin d'études chez ")
set_run(p, 1, "La Royale Gaming Investment")
set_run(p, 2, ", à Lisbonne, dans le secteur du jeu en ligne, m'a permis de construire des pipelines dbt/Dataform sur BigQuery à fort volume — des enjeux de fiabilité et de performance assez proches, je crois, de ceux d'une plateforme de billetterie et de marketing. J'y ai aussi optimisé la performance et les coûts des traitements ETL de 15 à 20 %, et créé des dashboards Power BI utilisés par les équipes métier.")

# --- Para 18 : compétences techniques (honnête sur les manques) ---
p = paras[18]
set_run(p, 0, "Sur le plan technique, je pratique Python et SQL au quotidien, ainsi que Snowflake et Power BI/Tableau pour la restitution. Je n'ai pas encore utilisé Dagster en production, mais j'ai commencé à monter en compétence sur Airflow de mon côté, et mon expérience sur dbt/Dataform m'a montré que je m'approprie vite un nouvel orchestrateur.")

# --- Para 19 : savoir-être ---
p = paras[19]
set_run(p, 0, "Autonome et à l'aise dans les environnements où l'on porte plusieurs sujets à la fois, j'apprécie particulièrement l'idée de rejoindre une structure d'une vingtaine de personnes, où chaque contribution compte et où les décisions se prennent vite.")

# --- Para 20 : clôture ---
p = paras[20]
set_run(p, 0, "Je serais ravi d'échanger avec vous sur la manière dont mon profil peut contribuer au développement de votre plateforme data.")

# Para 21 (formule de politesse) : inchangée
# Para 26 (signature) : inchangée

# --- Nettoyage du pied de page (résidu d'une candidature précédente) ---
for s in d.sections:
    for p in s.footer.paragraphs:
        for r in p.runs:
            r.text = ""

d.save(OUT)
print("LM générée :", OUT)

# --- Vérification : ré-ouvre et affiche le texte final ---
d2 = docx.Document(OUT)
print("\n=== CONTENU FINAL ===")
for i, p in enumerate(d2.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")
print("\n=== FOOTER ===")
for s in d2.sections:
    for i, p in enumerate(s.footer.paragraphs):
        print(i, repr(p.text))
