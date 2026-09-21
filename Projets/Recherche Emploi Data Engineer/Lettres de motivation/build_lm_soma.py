#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Construit la lettre de motivation SOMA (Analytics Engineer, Paris) à partir du modèle.
Adresse : Île-de-France (règle du profil, offre à Paris)."""
import docx

SRC = "_MODELE_lettre_motivation.docx"
OUT = "LM_ROBALO_Victor_SOMA_AnalyticsEngineer.docx"

d = docx.Document(SRC)
paras = d.paragraphs

def set_run(p, i, text):
    p.runs[i].text = text

# --- Para 1 : date (pas de mention de ville, cf règle Île-de-France -> pas de rue) ---
p = paras[1]
set_run(p, 2, "")
set_run(p, 3, "07")
set_run(p, 4, "/0")
set_run(p, 5, "9")
set_run(p, 6, "/2026")

# --- Para 2 : adresse -> Mobile Île-de-France ---
p = paras[2]
set_run(p, 1, " Mobile Île-de-France")

# --- Para 3 : ligne code postal / ville -> vide ---
p = paras[3]
set_run(p, 0, "")

# --- Para 6 : nom entreprise (bandeau droit) ---
p = paras[6]
set_run(p, 3, "SOMA")

# --- Para 7 : intitulé du poste (bandeau droit) ---
p = paras[7]
set_run(p, 0, "\t               Analytics Engineer")

# --- Para 13 : objet ---
p = paras[13]
set_run(p, 0, "Objet : Candidature au poste d'Analytics Engineer")

# --- Para 15 : accroche ---
p = paras[15]
set_run(p, 0, "C'est avec un vif intérêt que je vous transmets ma candidature pour le poste d'Analytics Engineer. Transformer la donnée brute en décisions qui comptent est précisément ce qui m'anime dans mon métier : au-delà de la construction de pipelines fiables, j'aime aller jusqu'au bout, jusqu'au dashboard qui fait vraiment décider.")
set_run(p, 1, "")
set_run(p, 2, "")

# --- Para 16 : expérience Sanofi ---
p = paras[16]
set_run(p, 0, "Mon expérience de trois ans en alternance chez Sanofi m'a appris ce que signifie fiabiliser un flux de données en production : j'y ai transformé des données brutes en indicateurs de performance (KPIs) exploités par les équipes métier et le Comité de Direction, dans un environnement industriel exigeant sur la qualité et la traçabilité des données.")

# --- Para 17 : expérience La Royale Gaming Investment ---
p = paras[17]
set_run(p, 0, "Mon stage de fin d'études chez ")
set_run(p, 1, "La Royale Gaming Investment")
set_run(p, 2, ", à Lisbonne, m'a permis de construire et maintenir des pipelines ETL/ELT (dbt/Dataform) sur BigQuery, de la donnée brute jusqu'à des dashboards Power BI utilisés quotidiennement par les équipes métiers pour piloter leurs décisions. J'y ai aussi optimisé la performance et les coûts des traitements de 15 à 20 %, tout en accompagnant les utilisateurs finaux vers une meilleure autonomie sur leurs indicateurs.")

# --- Para 18 : compétences techniques ---
p = paras[18]
set_run(p, 0, "Sur le plan technique, je pratique Python et SQL avancé au quotidien, ainsi que Snowflake et BigQuery pour la modélisation et l'optimisation d'entrepôts de données, et Dataform (l'équivalent dbt sur BigQuery) pour l'industrialisation des transformations. Power BI et Tableau sont mes outils de restitution habituels pour rendre la donnée actionnable auprès de publics non techniques.")

# --- Para 19 : savoir-être ---
p = paras[19]
set_run(p, 0, "Pragmatique et autonome, j'apprécie les environnements où l'on va jusqu'au bout de la chaîne de valeur de la donnée, de la fiabilisation du pipeline jusqu'à la décision métier, plutôt que de se limiter à produire des rapports.")

# --- Para 20 : clôture ---
p = paras[20]
set_run(p, 0, "Je serais ravi d'échanger avec vous sur la manière dont mon profil peut contribuer à la conception et à l'industrialisation de vos solutions analytiques.")

# Para 21 (formule de politesse) : inchangée
# Para 26 (signature) : inchangée

# --- Nettoyage du pied de page (résidu d'une candidature précédente) ---
for s in d.sections:
    for p in s.footer.paragraphs:
        for r in p.runs:
            r.text = ""

d.save(OUT)
print("LM générée :", OUT)

d2 = docx.Document(OUT)
print("\n=== CONTENU FINAL ===")
for i, p in enumerate(d2.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")
print("\n=== FOOTER ===")
for s in d2.sections:
    for i, p in enumerate(s.footer.paragraphs):
        print(i, repr(p.text))
