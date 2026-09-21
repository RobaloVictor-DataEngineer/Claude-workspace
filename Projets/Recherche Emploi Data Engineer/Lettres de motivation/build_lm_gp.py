#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Construit la lettre de motivation Groupe Premium (Data Engineer, Bois-Guillaume/Rouen).
Adresse : Rouen, inchangée (règle du profil, offre à Rouen)."""
import docx

SRC = "_MODELE_lettre_motivation.docx"
OUT = "LM_ROBALO_Victor_GroupePremium_DataEngineer.docx"

d = docx.Document(SRC)
paras = d.paragraphs

def set_run(p, i, text):
    p.runs[i].text = text

# --- Para 1 : date (adresse Rouen inchangée) ---
p = paras[1]
set_run(p, 3, "07")
set_run(p, 4, "/0")
set_run(p, 5, "9")
set_run(p, 6, "/2026")

# Para 2, 3 : adresse Rouen -> inchangées

# --- Para 6 : nom entreprise (bandeau droit) ---
p = paras[6]
set_run(p, 3, "Groupe Premium")

# --- Para 7 : intitulé du poste (bandeau droit) ---
p = paras[7]
set_run(p, 0, "\t               Data Engineer")

# --- Para 13 : objet ---
p = paras[13]
set_run(p, 0, "Objet : Candidature au poste de Data Engineer")

# --- Para 15 : accroche ---
p = paras[15]
set_run(p, 0, "C'est avec un vif intérêt que je vous transmets ma candidature pour le poste de Data Engineer. Basé à Rouen, je serais ravi de rejoindre une équipe qui construit et industrialise une plateforme data au service du courtage en assurance, avec un impact direct sur la conformité réglementaire et le pilotage de l'activité.")
set_run(p, 1, "")
set_run(p, 2, "")

# --- Para 16 : expérience Sanofi ---
p = paras[16]
set_run(p, 0, "Mon expérience de trois ans en alternance chez Sanofi m'a appris ce que signifie fiabiliser un flux de données en production : j'y ai transformé des données brutes en indicateurs de performance (KPIs) exploités par les équipes métier et le Comité de Direction, dans un environnement industriel exigeant sur la qualité, la sécurité et la conformité des données.")

# --- Para 17 : expérience La Royale Gaming Investment ---
p = paras[17]
set_run(p, 0, "Mon stage de fin d'études chez ")
set_run(p, 1, "La Royale Gaming Investment")
set_run(p, 2, ", à Lisbonne, m'a permis de construire et maintenir des pipelines ETL (dbt/Dataform) sur BigQuery, avec des contrôles qualité sur les données et une modélisation orientée besoins décisionnels. J'y ai aussi optimisé la performance et les coûts des traitements de 15 à 20 %, en garantissant la fiabilité des flux au quotidien.")

# --- Para 18 : compétences techniques (honnête sur les manques + motivation) ---
p = paras[18]
set_run(p, 0, "Sur le plan technique, je pratique SQL au quotidien ainsi que SQL Server, Snowflake et BigQuery, avec de solides bases en modélisation de données dimensionnelle et relationnelle. Je n'ai en revanche pas encore d'expérience sur Azure, Microsoft Fabric ni PySpark. J'ai cependant l'habitude de monter rapidement en compétence sur de nouveaux outils : c'est ce que j'ai fait avec Snowflake et avec Dataform, appris en autonomie pour industrialiser des pipelines dbt-like sur BigQuery. Je suis très motivé pour appliquer cette même dynamique à l'écosystème Microsoft Fabric et à PySpark.")

# --- Para 19 : savoir-être ---
p = paras[19]
set_run(p, 0, "Pragmatique et rigoureux, j'apprécie le travail en lien étroit avec un Data Lead et des équipes IT, ainsi que les environnements où la qualité et la fiabilité des données sont une priorité partagée.")

# --- Para 20 : clôture ---
p = paras[20]
set_run(p, 0, "Je serais ravi d'échanger avec vous sur la manière dont mon profil et ma motivation à apprendre peuvent contribuer à la construction de votre plateforme data.")

# Para 21 (formule de politesse) : inchangée
# Para 26 (signature) : inchangée

# --- Nettoyage du pied de page (résidu d'une candidature précédente) ---
for s in d.sections:
    for p in s.footer.paragraphs:
        for r in p.runs:
            r.text = ""

d.save(OUT)
print("LM générée :", OUT)

d2 = docx.Document(OUT)
print("\n=== CONTENU FINAL ===")
for i, p in enumerate(d2.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")
print("\n=== FOOTER ===")
for s in d2.sections:
    for i, p in enumerate(s.footer.paragraphs):
        print(i, repr(p.text))
