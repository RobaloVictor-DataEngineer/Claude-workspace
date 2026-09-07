#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Construit la lettre de motivation Aneo (Développeur Python Senior - programme
actuariat) à partir du modèle, en ne touchant qu'au texte des runs concernés —
police/taille/gras/alignement du modèle sont conservés intégralement."""
import docx

SRC = "_MODELE_lettre_motivation.docx"
OUT = "LM_ROBALO_Victor_Aneo_DeveloppeurPythonSenior.docx"

d = docx.Document(SRC)
paras = d.paragraphs

def set_run(p, i, text):
    p.runs[i].text = text

# --- Para 1 : date ---
p = paras[1]
set_run(p, 2, "Rouen, ")
set_run(p, 3, "04")
set_run(p, 4, "/0")
set_run(p, 5, "9")
set_run(p, 6, "/2026")

# --- Para 6 : nom entreprise (bandeau droit) ---
p = paras[6]
set_run(p, 3, "Aneo")

# --- Para 7 : intitulé du poste (bandeau droit) ---
p = paras[7]
set_run(p, 0, "\t               Développeur Python")

# --- Para 13 : objet ---
p = paras[13]
set_run(p, 0, "Objet : Candidature au poste de Développeur Python F/H — programme de transformation actuarielle")

# --- Para 15 : accroche ---
p = paras[15]
set_run(p, 0, "C'est avec un vif intérêt que je vous adresse ma candidature pour rejoindre ")
set_run(p, 1, "Aneo")
set_run(p, 2, ", et contribuer, aux côtés de vos équipes, au programme de transformation actuarielle mené pour l'un de vos clients du secteur assurance. Le remplacement progressif d'outils historiques par une plateforme Python industrialisée est exactement le type de projet où j'ai envie de mettre mon expérience de data engineer à l'épreuve.")

# --- Para 16 : expérience Sanofi ---
p = paras[16]
set_run(p, 0, "Mon expérience de trois ans en alternance chez Sanofi m'a permis d'évoluer au cœur d'un environnement industriel complexe et fortement réglementé. J'y ai appris qu'une donnée n'a de valeur que si elle est fiable, documentée et traçable : j'ai piloté des flux de données de bout en bout, de la transformation de données brutes en indicateurs de performance (KPIs) jusqu'à la création de dashboards d'aide à la décision pour les équipes métier et le Comité de Direction.")

# --- Para 17 : expérience La Royale Gaming Investment ---
p = paras[17]
set_run(p, 0, "Mon stage de fin d'études chez ")
set_run(p, 1, "La Royale Gaming Investment")
set_run(p, 2, ", à Lisbonne, a complété ce profil par une dimension plus opérationnelle : j'y ai optimisé la performance et les coûts de traitements ETL sur BigQuery (−15 à −20 %), et échangé régulièrement avec les équipes métier pour traduire leurs besoins en solutions concrètes — un exercice de traduction technique/métier assez proche, je crois, de celui que je retrouverais aux côtés de vos actuaires.")

# --- Para 18 : compétences techniques (honnête sur les manques) ---
p = paras[18]
set_run(p, 0, "Sur le plan technique, je pratique Python et SQL au quotidien, ainsi que la construction de pipelines de données sur BigQuery et dbt/Dataform ; j'ai également développé de mon côté un pipeline ETL complet (extraction, transformation avec pandas, chargement via SQLAlchemy). Databricks, Dataiku et PySpark ne font pas encore partie de mes outils du quotidien, mais mon parcours m'a montré que je monte vite en compétence sur une nouvelle stack — c'est ce que j'ai fait avec dbt et Dataform en quelques semaines lors de mon stage.")

# --- Para 19 : savoir-être ---
p = paras[19]
set_run(p, 0, "Pragmatique et orienté solutions, je ne cherche pas seulement à écrire du code : je cherche à comprendre le besoin — ici celui de vos actuaires — pour construire des traitements robustes et justes. J'apprécie aussi, plus largement, la culture d'autonomie et l'investissement d'Aneo dans la formation, qui comptent beaucoup pour moi dans le choix d'une entreprise.")

# --- Para 20 : clôture ---
p = paras[20]
set_run(p, 0, "Je serais ravi d'échanger avec vous sur la manière dont mon profil peut s'intégrer à l'équipe et contribuer à ce programme.")

# Para 21 (formule de politesse) : inchangée
# Para 26 (signature) : inchangée

# --- Nettoyage du pied de page (résidu d'une candidature précédente) ---
for s in d.sections:
    for p in s.footer.paragraphs:
        for r in p.runs:
            r.text = ""

d.save(OUT)
print("LM générée :", OUT)

# --- Vérification : ré-ouvre et affiche le texte final ---
d2 = docx.Document(OUT)
print("\n=== CONTENU FINAL ===")
for i, p in enumerate(d2.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")
print("\n=== FOOTER ===")
for s in d2.sections:
    for i, p in enumerate(s.footer.paragraphs):
        print(i, repr(p.text))
