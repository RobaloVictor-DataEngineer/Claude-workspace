#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Construit la lettre de motivation EKINOX (Data Engineer, Lille) à partir du modèle,
en ne touchant qu'au texte des runs concernés — police/taille/gras/alignement du
modèle sont conservés intégralement. Adresse : Lille (règle du profil)."""
import docx

SRC = "_MODELE_lettre_motivation.docx"
OUT = "LM_ROBALO_Victor_EKINOX_DataEngineer.docx"

d = docx.Document(SRC)
paras = d.paragraphs

def set_run(p, i, text):
    p.runs[i].text = text

# --- Para 1 : date (pas de mention de ville, cf règle Lille -> pas de rue) ---
p = paras[1]
set_run(p, 2, "")
set_run(p, 3, "07")
set_run(p, 4, "/0")
set_run(p, 5, "9")
set_run(p, 6, "/2026")

# --- Para 2 : adresse -> Mobile Lille ---
p = paras[2]
set_run(p, 1, " Mobile Lille")

# --- Para 3 : ligne code postal / ville -> vide ---
p = paras[3]
set_run(p, 0, "")

# --- Para 6 : nom entreprise (bandeau droit) ---
p = paras[6]
set_run(p, 3, "EKINOX")

# --- Para 7 : intitulé du poste (bandeau droit) ---
p = paras[7]
set_run(p, 0, "\t               Data Engineer")

# --- Para 13 : objet ---
p = paras[13]
set_run(p, 0, "Objet : Candidature au poste de Data Engineer")

# --- Para 15 : accroche ---
p = paras[15]
set_run(p, 0, "C'est avec un vif intérêt que je vous transmets ma candidature pour le poste de Data Engineer, au sein d'une entreprise qui construit une plateforme de « Data Sharing » pour estimer et prédire la valeur résiduelle de ses produits. Contribuer directement, aux côtés du CTO, à l'architecture des pipelines qui alimentent cet algorithme de prédiction est exactement le type de défi technique où j'ai envie d'apporter ma contribution.")
set_run(p, 1, "")
set_run(p, 2, "")

# --- Para 16 : expérience Sanofi ---
p = paras[16]
set_run(p, 0, "Mon expérience de trois ans en alternance chez Sanofi m'a appris ce que signifie fiabiliser un flux de données en production : j'y ai transformé des données brutes en indicateurs de performance (KPIs) exploités par les équipes métier et le Comité de Direction, dans un environnement industriel exigeant sur la qualité et la traçabilité des données.")

# --- Para 17 : expérience La Royale Gaming Investment ---
p = paras[17]
set_run(p, 0, "Mon stage de fin d'études chez ")
set_run(p, 1, "La Royale Gaming Investment")
set_run(p, 2, ", à Lisbonne, m'a permis de construire et maintenir des pipelines dbt/Dataform sur BigQuery à fort volume, en structurant les données brutes en couches exploitables pour la restitution — une logique assez proche, je crois, de l'organisation Bronze/Silver/Gold que vous mettez en place. J'y ai aussi optimisé la performance et les coûts des traitements ETL de 15 à 20 %, sur des flux alimentant des dashboards Power BI utilisés par les équipes métier.")

# --- Para 18 : compétences techniques (honnête sur les manques) ---
p = paras[18]
set_run(p, 0, "Sur le plan technique, je pratique Python et SQL au quotidien, ainsi que Snowflake, PostgreSQL et BigQuery pour la construction et l'optimisation de pipelines de données. Je n'ai en revanche pas encore d'expérience opérationnelle sur AWS ou Azure : mon expérience cloud s'est construite via BigQuery (Google Cloud), mais je suis motivé pour monter rapidement en compétence sur un nouvel écosystème cloud, comme je l'ai fait avec Snowflake.")

# --- Para 19 : savoir-être ---
p = paras[19]
set_run(p, 0, "Pragmatique et autonome, j'apprécie particulièrement les environnements où l'on peut échanger directement avec les décideurs techniques et où la simplicité de maintenance prime sur la complexité — c'est ce que je recherche dans une structure à taille humaine comme la vôtre.")

# --- Para 20 : clôture ---
p = paras[20]
set_run(p, 0, "Je serais ravi d'échanger avec vous sur la manière dont mon profil peut contribuer à la construction de votre plateforme de Data Sharing.")

# Para 21 (formule de politesse) : inchangée
# Para 26 (signature) : inchangée

# --- Nettoyage du pied de page (résidu d'une candidature précédente) ---
for s in d.sections:
    for p in s.footer.paragraphs:
        for r in p.runs:
            r.text = ""

d.save(OUT)
print("LM générée :", OUT)

d2 = docx.Document(OUT)
print("\n=== CONTENU FINAL ===")
for i, p in enumerate(d2.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")
print("\n=== FOOTER ===")
for s in d2.sections:
    for i, p in enumerate(s.footer.paragraphs):
        print(i, repr(p.text))
